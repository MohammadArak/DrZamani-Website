"""Authorization, data minimization, migration-compatible roles and owner protection."""

import ast
import secrets
from concurrent.futures import ThreadPoolExecutor
from datetime import date, time
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, select
from starlette.websockets import WebSocketDisconnect

from app.access import CATALOG, CODES, identity, seed_access
from app.browser_sessions import cookie_name, csrf_token, staff_state_hash
from app.database import Base, SessionLocal, get_engine
from app.main import app
from app.models import (
    AuthRateLimit,
    AuthSession,
    ConsultationMessage,
    Patient,
    Permission,
    Role,
    Service,
    StaffUser,
    Appointment,
)
from app.security import hash_password, issue_session_token, staff_session_expiry

PREFIX = "/api/v1/staff"
PASSWORD = "RoleFixturePassword123!"


@pytest.fixture(autouse=True)
def schema():
    Base.metadata.create_all(get_engine())
    with SessionLocal.begin() as db:
        seed_access(db)
        db.execute(delete(AuthRateLimit))


@pytest.fixture
def client():
    with TestClient(app) as value:
        yield value


def account(codes=None, *, builtin=None, active=True):
    with SessionLocal.begin() as db:
        if builtin:
            role = db.scalar(select(Role).where(Role.slug == builtin))
        else:
            role = Role(
                slug=f"test-{secrets.token_hex(8)}",
                name="نقش آزمایشی",
                is_active=True,
                permissions=list(
                    db.scalars(
                        select(Permission).where(Permission.code.in_(codes or []))
                    )
                ),
            )
            db.add(role)
        staff = StaffUser(
            username=f"role-{secrets.token_hex(8)}",
            full_name="کاربر آزمایشی",
            role="custom",
            password_hash=hash_password(PASSWORD),
            is_active=active,
            roles=[role],
        )
        db.add(staff)
        db.flush()
        token, digest = issue_session_token()
        db.add(
            AuthSession(
                token_hash=digest,
                staff_id=staff.id,
                expires_at=staff_session_expiry(),
                staff_state_hash=staff_state_hash(staff),
            )
        )
        return staff.id, role.id, token


def headers(token):
    return {"Authorization": f"Bearer {token}"}


def owner():
    return account(builtin="superadmin")


def login(client, staff_id):
    with SessionLocal() as db:
        username = db.get(StaffUser, staff_id).username
    captcha = client.get(PREFIX + "/auth/captcha").json()
    result = client.post(
        PREFIX + "/auth/login",
        json=dict(
            username=username,
            password=PASSWORD,
            captcha_id=captcha["captcha_id"],
            captcha_answer=captcha["debug_answer"],
        ),
    )
    assert result.status_code == 200, result.text
    return result.json()


READS = [
    "/dashboard",
    "/patients",
    "/patients/999999",
    "/settings",
    "/schedule",
    "/exceptions",
    "/services",
    "/services/999999/exceptions",
    "/appointments",
    "/appointments/export",
    "/consultations",
    "/appointments/999999/consultation",
    "/appointments/999999/consultation/images/999999",
    "/waitlist",
    "/finance/summary",
    "/finance/payments",
    "/audit-logs",
    "/sms/rules",
    "/sms/outbox",
    "/sms/campaigns",
    "/access/roles",
    "/access/permissions",
    "/access/staff",
]


@pytest.mark.parametrize("path", READS)
def test_empty_role_cannot_read_any_staff_data(client, path):
    _, _, token = account()
    assert client.get(PREFIX + path, headers=headers(token)).status_code == 403


WRITES = [
    ("PUT", "/settings", {}),
    ("POST", "/operations/run", {}),
    ("PUT", "/schedule", []),
    ("POST", "/exceptions", {}),
    ("DELETE", "/exceptions/999999", None),
    ("POST", "/services/999999/exceptions", {}),
    ("DELETE", "/services/999999/exceptions/999999", None),
    ("POST", "/services", {}),
    ("PUT", "/services/999999", {}),
    ("PATCH", "/patients/999999", {}),
    ("PATCH", "/appointments/999999", {"status": "confirmed"}),
    ("PATCH", "/appointments/999999/reschedule", {}),
    ("POST", "/appointments/999999/consultation/messages", {"body": "test"}),
    ("PATCH", "/waitlist/999999", {}),
    ("PATCH", "/finance/payments/999999/refund", {}),
    ("PUT", "/sms/rules/unknown", {}),
    ("POST", "/sms/campaigns/preview", {}),
    ("POST", "/sms/campaigns", {}),
    ("POST", "/sms/dispatch", {}),
    ("POST", "/access/roles", {"name": "Denied"}),
    ("PUT", "/access/roles/999999", {"name": "Denied"}),
    ("DELETE", "/access/roles/999999", None),
    ("POST", "/access/staff", {}),
    ("PUT", "/access/staff/999999", {}),
]


@pytest.mark.parametrize("method,path,payload", WRITES)
def test_empty_role_cannot_mutate(client, method, path, payload):
    _, _, token = account()
    assert (
        client.request(
            method, PREFIX + path, headers=headers(token), json=payload
        ).status_code
        == 403
    )


def test_every_staff_route_declares_permission_and_catalog_is_consistent():
    module = ast.parse(
        (Path(__file__).parents[1] / "app/routers/staff.py").read_text(encoding="utf-8")
    )
    for node in module.body:
        if not isinstance(node, ast.FunctionDef) or not node.decorator_list:
            continue
        dependencies = [
            n
            for n in ast.walk(node.args)
            if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)
        ]
        if node.name == "me":
            continue
        permissions = [
            n.args[0].value for n in dependencies if n.func.id == "require_permission"
        ]
        assert len(permissions) == 1, node.name
        assert permissions[0] in CODES
    assert len(CATALOG) == len(CODES)
    assert all(set(required) <= CODES for _, _, _, required, _ in CATALOG)


def test_accountant_can_see_finance_without_clinical_or_owner_access(client):
    staff_id, _, token = account(builtin="accountant")
    profile = client.get(PREFIX + "/me", headers=headers(token)).json()
    assert profile["role_titles"] == ["حسابدار"] and not profile["is_superadmin"]
    for path in ["/finance/summary", "/finance/payments", "/dashboard"]:
        assert client.get(PREFIX + path, headers=headers(token)).status_code == 200
    stats = client.get(PREFIX + "/dashboard", headers=headers(token)).json()
    assert (
        stats["patients_total"]
        == stats["today_total"]
        == stats["unread_conversations"]
        == 0
    )
    assert stats["daily_appointments"] == []
    for path in [
        "/patients",
        "/appointments",
        "/consultations",
        "/sms/outbox",
        "/access/roles",
    ]:
        assert client.get(PREFIX + path, headers=headers(token)).status_code == 403
    assert login(client, staff_id)["permissions"] == profile["permissions"]
    with client.websocket_connect("/api/v1/realtime") as ws:
        ws.send_json({"token": token})
        with pytest.raises(WebSocketDisconnect):
            ws.receive_json()
    # Denied chat does not log the accountant out of the financial portal.
    assert client.get(PREFIX + "/me", headers=headers(token)).status_code == 200


def clinical_fixture():
    with SessionLocal.begin() as db:
        service = db.scalar(select(Service).where(Service.allows_media_chat.is_(True)))
        patient = Patient(
            phone=f"+989{secrets.randbelow(10**9):09d}",
            first_name="Synthetic",
            last_name="Patient",
            internal_note="PRIVATE-RECORD",
            tags_json='["PRIVATE-TAG"]',
            needs_follow_up=True,
        )
        appointment = Appointment(
            patient=patient,
            service=service,
            tracking_code=secrets.token_hex(8),
            appointment_date=date.today(),
            start_time=time(10),
            end_time=time(10, 30),
            status="pending",
            patient_note="PRIVATE-PATIENT",
            staff_note="PRIVATE-STAFF",
        )
        db.add_all([patient, appointment])
        db.flush()
        message = ConsultationMessage(
            appointment_id=appointment.id,
            sender_type="patient",
            body="PRIVATE-CHAT",
            stored_file_name=f"{secrets.token_hex(12)}.webp",
            original_file_name="PRIVATE-FILENAME.jpg",
        )
        db.add(message)
        db.flush()
        return patient.id, appointment.id, message.id


def test_appointments_and_excel_redact_notes_and_intake_independently(client):
    _, appointment_id, _ = clinical_fixture()
    _, _, token = account(
        {"appointments.view", "appointments.export", "appointments.edit"}
    )
    payload = client.get(PREFIX + "/appointments", headers=headers(token)).json()
    appt = next(a for a in payload["items"] if a["id"] == appointment_id)
    assert appt["patient_note"] is appt["staff_note"] is None
    assert appt["intake_submission"] is None and appt["intake_form"]["questions"] == []
    assert appt["image_requirements"] == [] and appt["consultation_enabled"] is False
    export = client.get(PREFIX + "/appointments/export", headers=headers(token))
    assert export.status_code == 200
    sheet = load_workbook(BytesIO(export.content)).active
    row = next(
        row
        for row in sheet.iter_rows(min_row=4, values_only=True)
        if row[1] == appt["tracking_code"]
    )
    assert row[10:13] == (None, None, None)
    assert (
        client.patch(
            PREFIX + f"/appointments/{appointment_id}",
            headers=headers(token),
            json={"status": "cancelled"},
        ).status_code
        == 403
    )
    assert (
        client.patch(
            PREFIX + f"/appointments/{appointment_id}",
            headers=headers(token),
            json={"status": "confirmed", "staff_note": "malicious"},
        ).status_code
        == 403
    )
    success = client.patch(
        PREFIX + f"/appointments/{appointment_id}",
        headers=headers(token),
        json={"status": "confirmed"},
    )
    assert success.status_code == 200 and success.json()["staff_note"] is None
    with SessionLocal() as db:
        assert db.get(Appointment, appointment_id).staff_note == "PRIVATE-STAFF"


def test_directory_does_not_leak_clinical_tags_and_record_requires_permission(client):
    patient_id, _, _ = clinical_fixture()
    _, _, token = account({"patients.view"})
    payload = client.get(PREFIX + "/patients", headers=headers(token)).json()
    patient = next(p for p in payload["items"] if p["id"] == patient_id)
    assert (
        patient["tags"] == []
        and patient["needs_follow_up"] is False
        and patient["appointment_count"] == 0
    )
    assert payload["available_tags"] == []
    assert (
        client.get(
            PREFIX + f"/patients/{patient_id}", headers=headers(token)
        ).status_code
        == 403
    )
    assert (
        client.get(
            PREFIX + "/patients?tag=PRIVATE-TAG", headers=headers(token)
        ).status_code
        == 403
    )
    _, _, record_token = account({"patients.view", "patients.records.view"})
    record = client.get(
        PREFIX + f"/patients/{patient_id}", headers=headers(record_token)
    ).json()
    assert record["internal_note"] == "PRIVATE-RECORD"
    assert (
        record["appointments"]
        == record["conversations"]
        == record["payments"]
        == record["timeline"]
        == []
    )


def test_chat_reader_without_image_or_send_permission(client):
    _, appointment_id, message_id = clinical_fixture()
    _, _, token = account({"appointments.view", "consultations.view"})
    response = client.get(
        PREFIX + f"/appointments/{appointment_id}/consultation", headers=headers(token)
    )
    assert response.status_code == 200 and response.json()[0]["body"] == "PRIVATE-CHAT"
    assert (
        response.json()[0]["has_image"] is False
        and response.json()[0]["original_file_name"] is None
    )
    assert (
        client.get(
            PREFIX + f"/appointments/{appointment_id}/consultation/images/{message_id}",
            headers=headers(token),
        ).status_code
        == 403
    )
    assert (
        client.post(
            PREFIX + f"/appointments/{appointment_id}/consultation/messages",
            headers=headers(token),
            json={"body": "denied"},
        ).status_code
        == 403
    )


def test_owner_role_crud_validation_and_no_escalation(client):
    _, _, token = owner()
    h = headers(token)
    assert client.get(PREFIX + "/access/permissions", headers=h).status_code == 200
    for codes in [["unknown"], ["roles.manage"], ["appointments.edit"], ["schedule.view"]]:
        assert (
            client.post(
                PREFIX + "/access/roles",
                headers=h,
                json={"name": "Invalid", "permissions": codes},
            ).status_code
            == 422
        )
    created = client.post(
        PREFIX + "/access/roles",
        headers=h,
        json={"name": "Viewer", "permissions": ["services.view"]},
    )
    assert created.status_code == 201
    role = created.json()
    assert (
        client.put(
            PREFIX + f"/access/roles/{role['id']}",
            headers=h,
            json={"name": "Changed", "permissions": [], "is_active": False},
        ).status_code
        == 200
    )
    assert (
        client.delete(PREFIX + f"/access/roles/{role['id']}", headers=h).status_code
        == 200
    )
    with SessionLocal() as db:
        root_id = db.scalar(select(Role.id).where(Role.slug == "superadmin"))
    assert (
        client.put(
            PREFIX + f"/access/roles/{root_id}", headers=h, json={"name": "Tamper"}
        ).status_code
        == 409
    )
    assert (
        client.delete(PREFIX + f"/access/roles/{root_id}", headers=h).status_code == 409
    )
    _, _, regular_token = account(builtin="admin")
    assert (
        client.post(
            PREFIX + "/access/roles",
            headers=headers(regular_token),
            json={"name": "Owner"},
        ).status_code
        == 403
    )
    # The legacy string alone is never owner authority.
    with SessionLocal.begin() as db:
        item = StaffUser(
            username=f"legacy-{secrets.token_hex(5)}",
            full_name="Legacy",
            role="superadmin",
            password_hash=hash_password(PASSWORD),
        )
        db.add(item)
        db.flush()
        assert identity(item)["is_superadmin"] is False


def test_staff_creation_multi_roles_disabled_roles_and_password_rotation(client):
    _, _, token = owner()
    h = headers(token)
    with SessionLocal() as db:
        role_ids = list(
            db.scalars(select(Role.id).where(Role.slug.in_(["accountant", "author"])))
        )
    payload = dict(
        username=f"new-{secrets.token_hex(5)}",
        full_name="New Staff",
        password=PASSWORD,
        role_ids=role_ids,
    )
    result = client.post(PREFIX + "/access/staff", headers=h, json=payload)
    assert result.status_code == 201, result.text
    staff = result.json()
    assert set(staff["role_ids"]) == set(role_ids)
    assert (
        "finance.view" in staff["permissions"]
        and "articles.create" in staff["permissions"]
    )
    assert (
        "patients.records.view" not in staff["permissions"]
        and "password" not in result.text
    )
    assert (
        client.post(PREFIX + "/access/staff", headers=h, json=payload).status_code
        == 409
    )
    session = login(client, staff["id"])["access_token"]
    change = dict(
        full_name="Renamed",
        role_ids=role_ids,
        is_active=True,
        password="AnotherStrongPassword123!",
    )
    assert (
        client.put(
            PREFIX + f"/access/staff/{staff['id']}", headers=h, json=change
        ).status_code
        == 200
    )
    assert client.get(PREFIX + "/me", headers=headers(session)).status_code == 401
    assert (
        client.put(
            PREFIX + f"/access/staff/{staff['id']}",
            headers=h,
            json={**change, "role_ids": [999999]},
        ).status_code
        == 422
    )
    assert (
        client.put(
            PREFIX + f"/access/staff/{staff['id']}",
            headers=h,
            json={**change, "is_superadmin": True},
        ).status_code
        == 422
    )


def test_role_update_revokes_member_sessions_and_live_chat(client):
    _, _, owner_token = owner()
    staff_id, role_id, token = account({"appointments.view", "consultations.view"})
    with client.websocket_connect("/api/v1/realtime") as ws:
        ws.send_json({"token": token})
        assert ws.receive_json()["type"] == "realtime.ready"
        response = client.put(
            PREFIX + f"/access/roles/{role_id}",
            headers=headers(owner_token),
            json={"name": "Removed", "permissions": []},
        )
        assert response.status_code == 200, response.text
        ws.send_json({"type": "ping"})
        with pytest.raises(WebSocketDisconnect) as failure:
            ws.receive_json()
        assert failure.value.code == 4401
    assert client.get(PREFIX + "/me", headers=headers(token)).status_code == 401
    assert login(client, staff_id)["permissions"] == []
    assert (
        client.delete(
            PREFIX + f"/access/roles/{role_id}", headers=headers(owner_token)
        ).status_code
        == 409
    )


def test_cookie_access_management_requires_csrf(client):
    _, _, token = owner()
    client.cookies.set(cookie_name("staff"), token)
    assert client.get(PREFIX + "/access/roles").status_code == 200
    assert (
        client.post(PREFIX + "/access/roles", json={"name": "CSRF"}).status_code == 403
    )
    result = client.post(
        PREFIX + "/access/roles",
        headers={"X-CSRF-Token": csrf_token(token)},
        json={"name": "Valid"},
    )
    assert result.status_code == 201


def test_last_owner_and_concurrent_removal(client):
    # Isolate the owner invariant without deleting shared fixture accounts.
    with SessionLocal.begin() as db:
        originals = [
            (s.id, s.is_active)
            for s in db.scalars(select(StaffUser)).unique()
            if any(r.slug == "superadmin" for r in s.roles)
        ]
        for staff_id, _ in originals:
            db.get(StaffUser, staff_id).is_active = False
    first, root_role, token = owner()
    try:
        payload = dict(full_name="Owner", role_ids=[root_role], is_active=False)
        assert (
            client.put(
                PREFIX + f"/access/staff/{first}", headers=headers(token), json=payload
            ).status_code
            == 409
        )
        with SessionLocal() as db:
            secretary_role = db.scalar(select(Role.id).where(Role.slug == "secretary"))
        assert (
            client.put(
                PREFIX + f"/access/staff/{first}",
                headers=headers(token),
                json={**payload, "is_active": True, "role_ids": [secretary_role]},
            ).status_code
            == 409
        )
        second, _, second_token = owner()

        def disable(staff_id, session):
            return client.put(
                PREFIX + f"/access/staff/{staff_id}",
                headers=headers(session),
                json=payload,
            ).status_code

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(
                pool.map(
                    lambda item: disable(*item),
                    [(first, token), (second, second_token)],
                )
            )
        assert sorted(results) == [200, 409]
        with SessionLocal() as db:
            assert sum(db.get(StaffUser, i).is_active for i in [first, second]) == 1
    finally:
        with SessionLocal.begin() as db:
            for staff_id, active in originals:
                db.get(StaffUser, staff_id).is_active = active


def test_explicit_new_install_owner_cli(client, monkeypatch):
    from app import setup_owner

    username = f"install-{secrets.token_hex(6)}"
    monkeypatch.setattr(
        "sys.argv",
        ["setup_owner", "--username", username, "--name", "Installation Owner"],
    )
    monkeypatch.setattr("getpass.getpass", lambda _: PASSWORD)
    setup_owner.main()
    with SessionLocal() as db:
        staff = db.scalar(select(StaffUser).where(StaffUser.username == username))
        assert identity(staff)["is_superadmin"] and identity(staff)[
            "permissions"
        ] == sorted(CODES)


def test_direct_permission_change_invalidates_bound_session(client):
    staff_id, role_id, token = account({"services.view"})
    assert client.get(PREFIX + "/services", headers=headers(token)).status_code == 200
    with SessionLocal.begin() as db:
        db.get(Role, role_id).permissions = []
    assert client.get(PREFIX + "/services", headers=headers(token)).status_code == 401
    assert login(client, staff_id)["permissions"] == []


def test_seed_does_not_restore_edited_builtin_permissions(client):
    with SessionLocal.begin() as db:
        role = db.scalar(select(Role).where(Role.slug == "author"))
        originals = list(role.permissions)
        role.permissions = []
        db.flush()
        seed_access(db)
        assert role.permissions == []
        role.permissions = originals
