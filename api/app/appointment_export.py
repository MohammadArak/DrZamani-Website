from __future__ import annotations

from io import BytesIO

import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Appointment


STATUS_LABELS = {
    "pending": "در انتظار تأیید",
    "confirmed": "تأیید شده",
    "completed": "انجام شده",
    "cancelled": "لغو شده",
}


def _jalali_date(value) -> str:
    date = jdatetime.date.fromgregorian(date=value)
    return f"{date.year:04d}/{date.month:02d}/{date.day:02d}"


def _local_phone(value: str) -> str:
    return f"0{value[3:]}" if value.startswith("+98") else value


def build_appointments_workbook(
    items: list[Appointment],
    *,
    doctor_name: str,
    include_sensitive_notes: bool = True,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "نوبت‌ها"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A4"

    columns = [
        ("ردیف", 8),
        ("کد پیگیری", 18),
        ("نام بیمار", 24),
        ("شماره موبایل", 18),
        ("خدمت", 34),
        ("تاریخ نوبت", 16),
        ("ساعت شروع", 14),
        ("ساعت پایان", 14),
        ("مدت (دقیقه)", 15),
        ("وضعیت", 18),
        ("سابقه مراجعه", 16),
        ("توضیح بیمار", 38),
        ("یادداشت مطب", 38),
    ]
    last_column = get_column_letter(len(columns))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = f"گزارش نوبت‌های مطب {doctor_name}"
    sheet["A1"].font = Font(name="Tahoma", size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="293241")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = f"تعداد نوبت‌های این گزارش: {len(items)}"
    sheet["A2"].font = Font(name="Tahoma", size=10, color="5F6C7B")
    sheet["A2"].alignment = Alignment(horizontal="center")

    header_fill = PatternFill("solid", fgColor="D5AB64")
    thin = Side(style="thin", color="D9E0E7")
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=3, column=index, value=title)
        cell.font = Font(name="Tahoma", size=10, bold=True, color="293241")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[3].height = 27

    for row_index, item in enumerate(items, start=4):
        patient_name = " ".join(
            filter(None, [item.patient.first_name, item.patient.last_name])
        ) or "تکمیل‌نشده"
        values = [
            row_index - 3,
            item.tracking_code,
            patient_name,
            _local_phone(item.patient.phone),
            item.service.title,
            _jalali_date(item.appointment_date),
            item.start_time.strftime("%H:%M"),
            item.end_time.strftime("%H:%M"),
            item.service.duration_minutes,
            STATUS_LABELS.get(item.status, item.status),
            ("بله" if item.has_previous_visit else "خیر") if include_sensitive_notes else "",
            (item.patient_note or "") if include_sensitive_notes else "",
            (item.staff_note or "") if include_sensitive_notes else "",
        ]
        fill = PatternFill("solid", fgColor="F7F9FB" if row_index % 2 == 0 else "FFFFFF")
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if isinstance(value, str):
                # User-authored text must never become an executable Excel formula.
                cell.data_type = "s"
            cell.font = Font(name="Tahoma", size=9, color="293241")
            cell.fill = fill
            cell.border = Border(bottom=Side(style="hair", color="E5EAF0"))
            cell.alignment = Alignment(
                horizontal="right" if column_index in {3, 5, 12, 13} else "center",
                vertical="center",
                wrap_text=column_index in {5, 12, 13},
            )
        sheet.row_dimensions[row_index].height = 29

    sheet.auto_filter.ref = f"A3:{last_column}{max(3, len(items) + 3)}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
